"""Track classification: document vs structured vs analytical.

Routing must be deterministic and explainable -- no LLM, and every decision
carries reasons the UI can show.
"""

from __future__ import annotations

import csv
from pathlib import Path

import pytest

from knowledge.classification import Track, TrackClassifier


@pytest.fixture
def classifier(settings) -> TrackClassifier:
    return TrackClassifier(settings.classification)


def _csv(path: Path, rows: int, cols: tuple[str, ...] = ("a", "b")) -> Path:
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(cols)
        for i in range(rows):
            w.writerow([f"v{i}"] * len(cols))
    return path


def test_document_formats_take_the_document_track(classifier, tmp_path):
    for name in ("policy.pdf", "deck.pptx", "memo.docx", "page.html", "note.md",
                 "raw.txt"):
        p = tmp_path / name
        p.write_bytes(b"x")
        result = classifier.classify(p)
        assert result.track is Track.document, name
        assert result.reasons, "a decision must explain itself"
        assert result.requires_optin is False


def test_small_table_is_structured_knowledge(classifier, tmp_path):
    result = classifier.classify(_csv(tmp_path / "services.csv", rows=20))
    assert result.track is Track.structured
    assert result.requires_optin is False
    assert result.signals["extension"] == ".csv"


def test_row_heavy_csv_is_analytical(classifier, settings, tmp_path):
    rows = settings.classification.analytical.max_rows + 50
    result = classifier.classify(_csv(tmp_path / "telemetry.csv", rows=rows))
    assert result.track is Track.analytical
    assert result.requires_optin is True
    assert any("rows" in r for r in result.reasons)


def test_oversized_file_is_analytical(classifier, settings, tmp_path):
    settings.classification.analytical.max_bytes = 500
    p = tmp_path / "big.csv"
    _csv(p, rows=200)
    assert p.stat().st_size > 500
    result = classifier.classify(p)
    assert result.track is Track.analytical
    assert any("MB" in r for r in result.reasons)


def test_unsupported_extension_is_rejected(classifier, tmp_path):
    p = tmp_path / "archive.zip"
    p.write_bytes(b"PK")
    result = classifier.classify(p)
    assert result.track is Track.unsupported


def test_force_track_overrides_the_probe(classifier, settings, tmp_path):
    """'Add to Knowledge Base' promotes an analytical file without re-probing."""
    rows = settings.classification.analytical.max_rows + 10
    p = _csv(tmp_path / "huge.csv", rows=rows)
    assert classifier.classify(p).track is Track.analytical

    forced = classifier.classify(p, force_track=Track.structured)
    assert forced.track is Track.structured
    assert forced.requires_optin is False
    assert "explicitly routed" in forced.reasons[0]


def test_disabling_classification_restores_document_only_behaviour(
    settings, tmp_path
):
    settings.classification.enabled = False
    c = TrackClassifier(settings.classification)
    p = _csv(tmp_path / "anything.csv", rows=99999)
    assert c.classify(p).track is Track.document


def test_json_record_count_drives_the_decision(classifier, settings, tmp_path):
    small = tmp_path / "small.json"
    small.write_text('[{"a": 1}, {"a": 2}]', encoding="utf-8")
    assert classifier.classify(small).track is Track.structured

    cap = settings.classification.analytical.max_rows
    big = tmp_path / "big.json"
    big.write_text("[" + ",".join('{"a":1}' for _ in range(cap + 5)) + "]",
                   encoding="utf-8")
    assert classifier.classify(big).track is Track.analytical


def test_classification_is_serialisable_for_the_api(classifier, tmp_path):
    result = classifier.classify(_csv(tmp_path / "t.csv", rows=3))
    d = result.to_dict()
    assert d["track"] == "structured"
    assert isinstance(d["reasons"], list)
    assert "requires_optin" in d and "signals" in d


# ---- spreadsheets -----------------------------------------------------------
#
# openpyxl ships in the optional 'structured' extra, so these skip on a
# core-only install rather than failing the suite.

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

needs_openpyxl = pytest.mark.skipif(
    openpyxl is None, reason="spreadsheet probing needs openpyxl"
)


@needs_openpyxl
def test_formula_heavy_workbook_is_analytical(classifier, settings, tmp_path):
    """A model full of formulas is working material, not reference data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, settings.classification.analytical.max_formula_cells + 60):
        ws.cell(row=r, column=1, value=r)
        ws.cell(row=r, column=2, value=f"=A{r}*1.21")
    p = tmp_path / "financial-model.xlsx"
    wb.save(p)

    result = classifier.classify(p)
    assert result.track is Track.analytical
    assert any("formula" in r for r in result.reasons)
    assert result.signals["formula_cells"] > 0


@needs_openpyxl
def test_plain_workbook_is_structured(classifier, tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["vendor", "region"])
    for i in range(30):
        ws.append([f"vendor-{i}", "EU"])
    p = tmp_path / "vendors.xlsx"
    wb.save(p)

    result = classifier.classify(p)
    assert result.track is Track.structured
    assert result.signals["sheets"] == 1


@needs_openpyxl
def test_many_sheets_is_analytical(classifier, settings, tmp_path):
    wb = openpyxl.Workbook()
    for i in range(settings.classification.analytical.max_sheets + 2):
        ws = wb.create_sheet(f"s{i}")
        ws.append(["a"])
        ws.append([1])
    p = tmp_path / "workbook.xlsx"
    wb.save(p)

    result = classifier.classify(p)
    assert result.track is Track.analytical
    assert any("sheets" in r for r in result.reasons)
